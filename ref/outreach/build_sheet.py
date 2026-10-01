from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

H = ["#","Organization","Tier","Why it fits Tanit XR","Best contact (verified)","Name source","Email or form","Email source URL","Suggested ask","Language","Timing","Caveats"]

send = [
["Carthagina","Tunisia","Tunisian volunteer-run heritage NGO that documents heritage with new technology. Tunisian partner on the El Jem Tapestry (CyArk, StoryCenter, INP, launched Apr 8, 2026) and MedinaPedia with Wikimedia TN and ASM. Closest peer we have in Tunisia.","Emna Mizouni, Founding President","https://www.carthagina.org/","carthagina.heritage.project@gmail.com","https://www.carthagina.org/","Co-program: joint object scan day in the medina of Tunis, a Splats With Phones session for their members","English","Now","Site's 'partner with us' link also goes to emna.mizouni@gmail.com; using the project address."],
["DigiArt Living Lab (DALL) / CréaTec","Tunisia","Creative-tech living lab with a branch in Nabeul. Built Neapolis-X (VR/AR at Neapolis) with TICDCE in 2024. Same town as you and the Storm Harry captures.","Samia Chelbi, Founder & CEO","https://dall4all.org/","contact@dall4all.org","https://dall4all.org/","Co-program: co-hosted Splats With Phones session or object scan day in Nabeul","French or English","Now","Neapolis-X is their project; frame ours as object-level scans that add to it, not compete."],
["Wikimedia Tunisie (Wiki Loves Monuments Tunisia)","Tunisia","Volunteer group running Wiki Loves Monuments Tunisia, live Oct 1 to 31, 2026. Same values: volunteers, open licences.","Houssem Hommi, Board member, external relations (President: Houssem Eddine Boubakri)","https://meta.wikimedia.org/wiki/Wikimedia_Tunisie","wikimedia-tn@lists.wikimedia.org","https://meta.wikimedia.org/wiki/Wikimedia_Tunisie","Co-program: a joint photowalk plus phone-scanning session during WLM this month","English or French","This month (WLM ends Oct 31)","Mailing list, may reach all members and may need a subscription to post. Commons is CC BY-SA."],
["ICOMOS Tunisia (national committee)","International heritage","The local network of Tunisian heritage professionals. Natural host for an activity on the International Day for Monuments and Sites (Apr 18).","Faïka Béjaoui, President","https://www.icomos.org/national-committees/","icomos.comitetunisien@gmail.com","https://www.icomos.org/national-committees/","Demo + co-hosted scan day or talk for Apr 18, 2027","French","Now, for Apr 2027","Gmail address is the one ICOMOS publishes officially."],
["ICOM Tunisia (museums committee)","International heritage","Our scans are museum and site objects (stelae, statues, capitals, lamps). ICOM Tunisia is the museum professionals' network.","Nejib Sellaouti, Chair","https://icom.museum/en/committee/icom-tunisia/","icom.tunisie.2020@gmail.com","https://icom.museum/en/committee/icom-tunisia/","Workshop: phone-scanning session for museum staff, tie-in to International Museum Day (May 18, 2027)","French","Now, for May 2027","ICOM page is undated; chair may have changed. Draft will greet the committee, not him by name, unless you know he is current."],
["Institut français de Tunisie","Foreign institutes","Funding digital heritage right now: SAWN call (with INP, ASM) selected 10 digital-heritage projects in Sep 2026. Runs médiathèques in Tunis, Sousse, Sfax.","Louis Logodin, Attaché culturel","https://www.institutfrancais-tunisie.com/ift","webmestre@institutfrancais-tunisie.com (general) or form","https://www.institutfrancais-tunisie.com/contact","Workshop: a free Splats With Phones object scan day at the Tunis médiathèque; intro to SAWN laureates; ask if EUNIC Tunisia is bidding for the Cluster Fund (Oct 30)","French","Now","No direct culture email published; goes via webmestre@ addressed to M. Logodin."],
["Patrimoine 3000 (EU / Expertise France)","Foreign institutes","EU-funded €19M programme rehabilitating the Carthage Museum and Byrsa Hill, closed to the public until 2027. We already have Byrsa Hill statues in 3D.","No named staff (team page under update)","https://patrimoine3000.tn/contact/","contact@patrimoine3000.tn","https://patrimoine3000.tn/contact/","Data sharing + demo: offer our Byrsa and Carthage object models for public outreach while the museum is closed","French","Now","INP controls the collections; ask is only to share what we already have."],
["MarEA, Maritime Endangered Archaeology (Southampton / Ulster)","International heritage","Records threats to coastal heritage with Tunisia as a focus, funded to July 2027. Our Neapolis captures after Storm Harry are ground evidence they rarely get.","Dr Lucy Blue, Project lead","https://marea.soton.ac.uk/team/","marea@soton.ac.uk","https://marea.soton.ac.uk/contact/","Data sharing: offer the Neapolis post-storm captures and photos as a condition record","English","Now (before winter storm season)","Same database as EAMENA (Oxford); writing to MarEA only."],
["UNESCO World Heritage Centre, Dive into Heritage","International heritage","UNESCO's 3D-plus-stories platform, pilot in the Arab States, next phase in preparation. Best UNESCO match for object stories.","Thomas Rigauts, Project Officer (Dive into Heritage)","https://whc.unesco.org/en/whoswho/action=detail&order=264044","t.rigauts@unesco.org","https://whc.unesco.org/en/whoswho/action=detail&order=264044","Data sharing: offer object models and stories for the next phase","English","Now","You wrote ARC-WH about Dive into Heritage on Sep 27; this goes to the UNESCO officer directly. UNESCO has no Tunis office (Rabat covers it, no public email)."],
["UMD Michelle Smith Collaboratory for Visual Culture","Digital heritage & XR","College Park. Co-leads a community photogrammetry project scanning Piscataway Conoy objects at the Smithsonian (Aug 2026). Closest academic match to our model, and local to you.","Quint Gregory, Director","https://arthistory.umd.edu/directory/henry-gregory","quint@umd.edu","https://arthistory.umd.edu/directory/henry-gregory","Demo/workshop: a talk + co-hosted scan session with UMD students; invite to Oct 22","English","Now","You already have UMD ties (XR Club, Bella)."],
["UF Digital Worlds, Prof. Angelos Barmpoutis","Digital heritage & XR","Directs the NEH Digital Epigraphy and Archaeology project and its toolbox for 3D capture of inscriptions. Our Tophet stelae carry inscriptions.","Angelos Barmpoutis, Professor, Digital Worlds Institute","https://angelosbarmpoutis.com/","angelos@digitalworlds.ufl.edu","https://angelosbarmpoutis.com/","Co-program: run a few Tophet stelae scans through the Digital Epigraphy Toolbox; guest lecture or student project","English","Now","Existing contact of yours, so the draft is warmer."],
["ASOR Cultural Heritage Initiatives","Digital heritage & XR","Alexandria, VA. Trained Tunisian partners in documentation and photogrammetry (Manouba heritage lab, Carthagina, Scouts 'Tourath').","No named staff with current titles","https://www.asor.org/chi/updates/Tunisia","info@asor.org","https://www.asor.org/chi/","Talk + intros: present Tanit XR to CHI staff; invite to Oct 22","English","Now","Tunisia updates are 2022 to 2023; current funding unknown."],
["ISAMM, Institut Supérieur des Arts Multimédia de la Manouba","Tunisia","Public institute with 3D, video games and VR engineering tracks. Volunteer pool and workshop host.","Skander Belhaj, Directeur","https://uma.rnu.tn/fr/university-establishments/8?name=institut-superieur-des-arts-multimedia-de-la-manouba","direction@isamm.uma.tn","https://uma.rnu.tn/fr/university-establishments/8?name=institut-superieur-des-arts-multimedia-de-la-manouba","Workshop: Gaussian splats and phone photogrammetry masterclass for 3D/VR students","French","Next Tunisia trip","Name and email from the university's page (undated)."],
["ISMPT, Institut Supérieur des Métiers du Patrimoine de Tunis","Tunisia","The only university institute in Tunisia devoted to heritage trades; fieldwork-based degrees.","Youssef Lahbib, Directeur","https://utunis.rnu.tn/fr/university-establishments/7?name=institut-superieur-des-metiers-du-patrimoine-de-tunis","ISMPT@ismpt.rnu.tn","https://utunis.rnu.tn/fr/university-establishments/7?name=institut-superieur-des-metiers-du-patrimoine-de-tunis","Workshop: one-day phone-scanning field session for heritage students; free course after","French","Next Tunisia trip","General inbox."],
]

optional = [
["Istituto Italiano di Cultura di Tunisi","Foreign institutes","Italy is Tunisia's main archaeology partner; Tunisia Heritage Center (mosaics) at the Bardo signed Jan 2026; heritage days with INP Apr 2026.","Fabio Ruggirello, Direttore","https://italiana.esteri.it/italiana/sedi/istituto-italiano-di-cultura-di-tunisi/","segreteria.iictunisi@esteri.it","https://iictunisi.esteri.it/it/chi-siamo/contatti/","Talk/demo on community scanning of mosaics and Punic/Roman objects","French","Any","Mostly programmes Italian content; talk slot possible, money unlikely."],
["FLAHM Manouba, lab 'Régions et Ressources Patrimoniales'","Tunisia","Heritage research lab and doctoral school; academic partner for a methods paper or student fieldwork.","Lamia Ben Abid, head of lab LR99ES23","https://flm.rnu.tn/fra/pages/415/Structures-de-recherches","lamia.benabid@flah.uma.tn","https://flm.rnu.tn/fra/pages/415/Structures-de-recherches","Letter of support + talk to doctoral students","French","Any","Page undated."],
["DCX, Digital Cultural eXperience","Tunisia","Tunisian heritage VR studio; 'Immersive Stories' caravan with the Institut français through all 24 governorates; school clubs.","No full name published","https://dcx.studio/","midani@dcx.studio","https://dcx.studio/","Data sharing: our free models for their school clubs","English or French","Any","Commercial studio, may see us as competition."],
["Kamel Lazaar Foundation (B7L9)","Tunisia","Funds arts and heritage projects; B7L9 art centre runs workshops and talks.","No verified name","https://www.kamellazaarfoundation.org/node/191","contact-tunis@kamellazaarfoundation.org","https://www.kamellazaarfoundation.org/node/191","Talk/workshop at B7L9 on immersive art and heritage","English or French","Any","Leans contemporary art; pitch your artist side."],
["Scan the World (MyMiniFactory)","Digital heritage & XR","Open archive of community-scanned cultural objects; invites people to preserve their culture; has volunteer chapters.","No named staff","https://www.myminifactory.com/users/Scan%20The%20World","stw@myminifactory.com","https://www.myminifactory.com/users/Scan%20The%20World","Co-program: a Tanit XR collection or Tunisia chapter","English","Any","Their files are for 3D printing; check our licences and museum permissions first."],
["Cultural Emergency Response (CER)","International heritage","Rolling first-aid grants after disasters, incl. emergency documentation; Tunisia eligible.","No named staff","https://www.culturalemergency.org/programs/first-aid-to-cultural-heritage","firstaid@culturalemergency.org","https://www.culturalemergency.org/programs/first-aid-to-cultural-heritage","Short intro now; apply only if a storm exposes ruins again","English","Before winter","Applicant must live and work in Tunisia (a Tunisia-based volunteer lead would front it)."],
["INP, Institut National du Patrimoine","Tunisia","The authority over every site and museum. Also documented the Nabeul coast after Storm Harry.","Mounir Fantar, Director, Programming, Cooperation and Publication","https://www.inp2020.tn/en/inp_tunisie/administration/","inptunis4@gmail.com","https://www.inp2020.tn/contact/","Letter offering our object scans + a demo","Arabic or French (formal)","Best in person on your next trip","Sensitive: collaboration not confirmed. Better via Carthagina intro or a visit than a cold email."],
]

calendar = [
["Barakat Trust","Grant up to £10k for documentation/digitisation of Islamic art and heritage (Kairouan, medina mihrab, Zaghouan objects)","Opens Jan 1, 2027","https://barakat.org/grants/","applications@barakat.org"],
["Unity for Humanity","Grant for SDG projects built in Unity; needs a working demo of the VR museum","Reopens 2027 (third-party says Feb 20, 2027, unverified)","https://unity.com/humanity","Form only"],
["CyArk Heritage Amplified + Training Grant","Grant + free training cohort for a volunteer lead","Likely Nov 2026 to Jan 2027 (unannounced)","https://www.cyark.org/grants/heritageAmplified/","info@cyark.org"],
["CIPA Heritage Documentation symposium","Paper on volunteer phone scanning, Abu Dhabi, Nov 13 to 19, 2027","Call for papers not posted yet","https://www.cipaheritagedocumentation.org/activities/conferences/","info@cipaheritagedocumentation.org"],
["British Council Connections Through Culture","Up to £10k for UK to Tunisia creative tech/XR collaborations (needs UK + Tunisian registered partners)","Next round ~mid 2027","https://www.britishcouncil.tn/en/programmes/arts/connections-through-culture-grants","info@tn.britishcouncil.org"],
["Villa Salammbô residency (IFT, Sousse)","One-month residency; you can't apply alone from DC (needs a Tunisia-resident + France-resident duo, e.g. two volunteers)","Deadline Oct 18, 2026","https://www.institutfrancais-tunisie.com/candidatures/villa-salammbo-saison-7-residence-internationale-de-recherche-et-de-creation-de-lift/","meriam.benamor@ / maud.tronet@institutfrancais-tunisie.com"],
["World Monuments Watch","Nominate Neapolis (needs INP on board)","Next cycle ~early 2028","https://www.wmf.org/watch-nominate","watch@wmf.org"],
]

contacted = [
["CyArk + Open Heritage 3D","Emailed Sep 27 (object scans for Open Heritage 3D)","NEW ANGLE: CyArk built the El Jem Tapestry with INP and Carthagina (Apr 2026). A short follow-up offering our object models to sit alongside it."],
["AMVPPC","Emailed Apr 20, 2026 (meeting request)","NEW ANGLE: an AR/3D station at the next Nuit des Musées (18 museums, Mar 2026 edition). DG: Rabiaa Belfeguira, dg.amvppc@amvppc.tn."],
["U.S. Embassy Tunis","AFCP already contacted","NEW ANGLE: CPAIG grants ($25k to $150k, anti-trafficking of antiquities; 3D records help identify stolen objects). 2026 notes were due Feb 8; watch for early 2027. Unverified whether 3D documentation is eligible."],
["ARC-WH","Emailed Sep 27 (Dive into Heritage)","Nothing new; UNESCO's own officer is in the main list."],
["Smithsonian DPO 3D (Vince Rossi)","Existing thread (April meetings, Sep 27 empty draft)","Continue that thread yourself; not redrafted."],
["GWU CAI / IMES, Tunisian Embassy DC","Emailed today (Oct 22 event)","Skip."],
["South Mediterranean University","Emailed Mar 5, 2026","Skip."],
["Gerda Henkel, National Geographic, Niantic Spatial","Already contacted","Nothing new found."],
]

skipped = [
["ALIPH","Conflict and post-conflict areas only; Tunisia not eligible"],
["ICCROM / ATHAR","Aimed at government professionals; no open award cycle"],
["Getty Foundation / GCI","Invitation only; works in Tunisia only through INP"],
["Goethe-Institut Tunis","Film, dance, translation; no heritage or digital track"],
["EU Delegation (direct)","No open culture call; heritage money runs via Patrimoine 3000"],
["UNESCO Rabat office","Right office but no public email; slow"],
["Ministry of Cultural Affairs, Bardo, Carthage Museum","Too broad / no working site / closed for works; reached via Patrimoine 3000, ICOM, INP"],
["TICDCE (Ministry digital culture centre)","Strong fit but no verified email (official page blocked from the US). Worth a visit at the Cité de la Culture or a check from Tunisia"],
["Fondation Rambourg, B'chira Art Center","Closed / inactive"],
["Sketchfab Cultural Heritage","Sketchfab changed hands in 2026 and the museum programme no longer exists"],
["Polycam, Luma, XGRIDS, Meta, Snap, Epic MegaGrants","No nonprofit or heritage programme, or needs Unreal"],
["Big research labs (Getty, CNR-ISTI, Duke, UCLA, Brown, Factum, Zamani, Oxford IDA)","No Tunisia or community link; reply unlikely"],
["Arc/k Project","Closest peer model but looks dormant since 2021"],
]

wb = Workbook()
hdr = PatternFill("solid", fgColor="111518"); gold = Font(bold=True, color="FFCD05")
def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]:
        c.fill = hdr; c.font = gold; c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
ws = wb.active; ws.title = "Send now"
sheet(ws, H, [[i+1]+r for i,r in enumerate(send)], [4,30,16,50,32,30,32,30,40,12,16,40])
sheet(wb.create_sheet("Optional"), H, [[i+1]+r for i,r in enumerate(optional)], [4,30,16,50,32,30,32,30,40,12,16,40])
sheet(wb.create_sheet("Apply later"), ["Program","What","When","URL","Contact"], calendar, [32,60,30,50,36])
sheet(wb.create_sheet("Already contacted"), ["Organization","Status","New angle?"], contacted, [32,40,80])
sheet(wb.create_sheet("Skipped"), ["Organization","Why skipped"], skipped, [45,90])
wb.save("/Users/inessaid/Documents/tanitxr.org/ref/outreach/tanitxr-partner-outreach.xlsx")
print("ok")
